"""Importação de carteira a partir de planilha da corretora.

Lê .xlsx/.xlsm (openpyxl) e .csv, e devolve linhas normalizadas prontas para
`modules/carteira.py`. Não grava nada: quem decide o que entra é a rota, e só
depois que a pessoa viu a prévia.

Quatro decisões que sustentam o resto:

  - **O cabeçalho é casado por PADRÃO, não por nome exato.** Cada corretora
    exporta de um jeito: "Código", "Ticker", "Papel", "Produto"; "Quantidade",
    "Qtde", "Qtd. Disponível"; "Preço médio", "Preço Médio (R$)", "PM". Lista
    fechada de nomes quebra na primeira corretora nova. Mesmo desenho do
    `_mapear_colunas` de `atualizar_fundos_cvm.py`.

  - **O cabeçalho pode não estar na primeira linha.** Export de corretora
    costuma vir com título, CNPJ e data antes da tabela. Procuramos a linha que
    parece cabeçalho nas primeiras `LINHAS_PROCURA_CABECALHO`; exigir linha 1
    reprovaria quase todo arquivo real.

  - **Uma linha ruim não derruba o arquivo.** Cada linha vira "ok" ou "erro"
    com o motivo e o número da linha na planilha. Recusar as 40 posições
    porque a 12ª tem um traço no lugar do preço é o tipo de rigor que faz a
    pessoa desistir e digitar tudo à mão.

  - **Lê em modo streaming e com teto.** Arquivo enviado por usuário é entrada
    não confiável: `read_only=True` não materializa a planilha inteira na
    memória, e `MAXIMO_LINHAS` impede que uma aba com um milhão de linhas vire
    consumo de RAM do servidor.
"""

import csv
import io
import re
import unicodedata

from modules import carteira

# Teto do arquivo aceito. Carteira de pessoa física não passa disso nem de
# longe; acima é engano ou abuso.
MAXIMO_BYTES = 2 * 1024 * 1024
MAXIMO_LINHAS = 500
LINHAS_PROCURA_CABECALHO = 25

EXTENSOES = (".xlsx", ".xlsm", ".csv", ".txt")

# Padrões sobre o nome NORMALIZADO da coluna (minúsculo, sem acento).
PADROES = {
    "ticker": [r"^(codigo|cod|ticker|papel|ativo|produto)", r"codigo.*negoc"],
    "quantidade": [r"^(quantidade|qtde|qtd|quant)", r"quantidade.*(disponivel|total|final)"],
    "preco_medio": [r"pre[cç]o.*m[eé]d", r"^(pm|preco medio|custo medio)", r"^valor.*m[eé]dio"],
}


class ErroImportacao(ValueError):
    """Arquivo inutilizável como um todo (formato, cabeçalho, tamanho)."""


def _normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = texto.replace("_", " ").replace("-", " ").replace(".", " ")
    return " ".join(texto.lower().split())


def mapear_colunas(cabecalho):
    """{campo: índice}. Primeiro padrão que casar vence; coluna já usada não
    é reaproveitada — "Preço" e "Preço médio" na mesma planilha não podem cair
    os dois no mesmo campo."""
    mapa, usados = {}, set()
    normalizados = [_normalizar(c) for c in cabecalho]
    for campo, padroes in PADROES.items():
        for padrao in padroes:
            for indice, texto in enumerate(normalizados):
                if indice in usados or not texto:
                    continue
                if re.search(padrao, texto):
                    mapa[campo] = indice
                    usados.add(indice)
                    break
            if campo in mapa:
                break
    return mapa


def _numero_br(bruto):
    """Aceita 1.234,56 (planilha brasileira) e 1234.56 (exportação em inglês).

    A ambiguidade real é "1.234": pode ser mil duzentos e trinta e quatro com
    separador de milhar, ou 1,234 em inglês. Resolvemos pelo formato: ponto
    seguido de exatamente 3 dígitos e sem vírgula na string é separador de
    milhar. Não é infalível, e por isso o valor aparece na prévia antes de
    virar posição.
    """
    if bruto is None:
        return None
    if isinstance(bruto, (int, float)):
        valor = float(bruto)
        return valor if valor == valor else None

    texto = str(bruto).strip()
    if not texto:
        return None
    texto = re.sub(r"[R$\s ]", "", texto, flags=re.IGNORECASE)
    if not texto or texto in ("-", "--", "—"):
        return None

    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", texto):
        texto = texto.replace(".", "")

    try:
        valor = float(texto)
    except ValueError:
        return None
    return valor if valor == valor else None


def _limpar_ticker(bruto):
    """Extrai o código de uma célula que pode trazer mais coisa junto.

    Export de corretora escreve "PETR4 - PETROBRAS PN" ou "PETR4F"
    (fracionário). O F do fracionário é removido: é o mesmo papel, e o
    investidor não pensa nele como ativo separado.
    """
    texto = str(bruto or "").strip().upper()
    if not texto:
        return ""
    casou = re.search(r"\b([A-Z]{4}\d{1,2})F?\b", texto)
    if casou:
        return casou.group(1)

    # Sem nada com forma de ticker: devolvemos a primeira palavra SÓ se ela
    # tiver algum dígito. Ticker sempre termina em número, então "PETRO4"
    # (erro de digitação) volta e é sinalizado, enquanto "TOTAL" e "Subtotal"
    # — rodapé de planilha — voltam vazios e são ignorados em silêncio.
    # Sinalizar rodapé como erro encheria a prévia de vermelho sem motivo.
    primeira = texto.split()[0] if texto.split() else ""
    return primeira if any(c.isdigit() for c in primeira) else ""


def _linhas_do_xlsx(conteudo):
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover
        raise ErroImportacao(
            "Leitura de Excel indisponível no servidor (openpyxl ausente). "
            "Exporte a planilha como CSV.")
    try:
        # read_only: streaming, não materializa a planilha inteira.
        # data_only: queremos o VALOR calculado, não a fórmula.
        livro = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as erro:  # noqa: BLE001
        raise ErroImportacao(f"Não consegui abrir a planilha: {type(erro).__name__}.")

    try:
        aba = livro[livro.sheetnames[0]]
        linhas = []
        for indice, linha in enumerate(aba.iter_rows(values_only=True), start=1):
            linhas.append(list(linha))
            if indice >= MAXIMO_LINHAS + LINHAS_PROCURA_CABECALHO:
                break
        return linhas
    finally:
        livro.close()


def _linhas_do_csv(conteudo):
    for codificacao in ("utf-8-sig", "latin-1"):
        try:
            texto = conteudo.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover
        raise ErroImportacao("Não consegui ler o arquivo (codificação desconhecida).")

    amostra = texto[:4000]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=";,\t")
        separador = dialeto.delimiter
    except csv.Error:
        # CSV brasileiro usa ";" porque a vírgula é decimal. Na dúvida, o que
        # aparecer mais na amostra.
        separador = ";" if amostra.count(";") >= amostra.count(",") else ","

    return [linha for linha in csv.reader(io.StringIO(texto), delimiter=separador)][
        :MAXIMO_LINHAS + LINHAS_PROCURA_CABECALHO]


def _achar_cabecalho(linhas):
    """(índice, mapa). O cabeçalho é a primeira linha em que ticker e
    quantidade aparecem juntos — só "ticker" casaria com o título do arquivo."""
    for indice, linha in enumerate(linhas[:LINHAS_PROCURA_CABECALHO]):
        mapa = mapear_colunas(linha)
        if "ticker" in mapa and "quantidade" in mapa:
            return indice, mapa
    return None, {}


def ler(conteudo, nome_arquivo):
    """[{linha, ticker, quantidade, preco_medio, erro}]. Levanta
    `ErroImportacao` só quando o arquivo inteiro é inutilizável."""
    if not conteudo:
        raise ErroImportacao("Arquivo vazio.")
    if len(conteudo) > MAXIMO_BYTES:
        raise ErroImportacao(
            f"Arquivo acima de {MAXIMO_BYTES // (1024 * 1024)} MB.")

    nome = (nome_arquivo or "").lower()
    if not nome.endswith(EXTENSOES):
        raise ErroImportacao(
            "Formato não aceito. Envie .xlsx, .xlsm ou .csv.")

    brutas = (_linhas_do_csv(conteudo) if nome.endswith((".csv", ".txt"))
              else _linhas_do_xlsx(conteudo))
    if not brutas:
        raise ErroImportacao("A planilha não tem linhas.")

    inicio, mapa = _achar_cabecalho(brutas)
    if inicio is None:
        raise ErroImportacao(
            "Não achei as colunas de papel e quantidade. O cabeçalho precisa "
            "ter uma coluna de código (Ticker, Papel, Código) e uma de "
            "quantidade — veja a planilha modelo.")

    tem_preco = "preco_medio" in mapa
    saida = []
    for deslocamento, linha in enumerate(brutas[inicio + 1:], start=inicio + 2):
        if len(saida) >= MAXIMO_LINHAS:
            break
        if not any(str(c or "").strip() for c in linha):
            continue

        def celula(campo):
            indice = mapa.get(campo)
            return linha[indice] if indice is not None and indice < len(linha) else None

        ticker = _limpar_ticker(celula("ticker"))
        quantidade = _numero_br(celula("quantidade"))
        preco = _numero_br(celula("preco_medio")) if tem_preco else None

        item = {"linha": deslocamento, "ticker": ticker,
                "quantidade": quantidade, "preco_medio": preco, "erro": None}

        if not ticker:
            continue  # linha de rodapé/total: ignorada em silêncio, não é erro
        if not carteira.FORMA_TICKER.match(ticker):
            item["erro"] = f"'{ticker}' não tem forma de ticker da B3."
        elif quantidade is None:
            item["erro"] = "Quantidade ausente ou ilegível."
        elif quantidade <= 0:
            item["erro"] = "Quantidade precisa ser maior que zero."
        elif not tem_preco:
            item["erro"] = ("Planilha sem coluna de preço médio — informe o "
                            "preço na tela para este papel.")
        elif preco is None:
            item["erro"] = "Preço médio ausente ou ilegível."
        elif preco <= 0:
            item["erro"] = "Preço médio precisa ser maior que zero."

        saida.append(item)

    if not saida:
        raise ErroImportacao("Achei o cabeçalho, mas nenhuma linha com dados.")
    return saida


def planilha_modelo():
    """.xlsx de exemplo, em bytes. Mostrar o formato aceito evita metade dos
    erros de importação antes de eles acontecerem."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    livro = Workbook()
    aba = livro.active
    aba.title = "Carteira"
    aba.append(["Ticker", "Quantidade", "Preço médio"])
    for celula in aba[1]:
        celula.font = Font(bold=True)
    for exemplo in (("PETR4", 100, 32.50), ("TAEE11", 200, 34.10),
                    ("HGLG11", 50, 158.00), ("CPLE6", 300, 11.20)):
        aba.append(list(exemplo))
    aba.column_dimensions["A"].width = 14
    aba.column_dimensions["B"].width = 14
    aba.column_dimensions["C"].width = 16

    buffer = io.BytesIO()
    livro.save(buffer)
    return buffer.getvalue()
